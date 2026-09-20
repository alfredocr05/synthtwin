"""Round 2 of the review, the disclosure pass: each test from its reproduction.

Seven items were raised against commit `05e7d89` by a reviewer reading
only for what a description GIVES AWAY about one person. Each is
reproduced here from the reviewer's own inputs, with the numbers it was
reproduced with written into the test that closes it, and each carries a
RED CHECK that withdraws the repair and shows the defect come back.

THE SEVEN, in the reviewer's order:

1. the BLOCKER -- one recognised spelling of no value in the readings
   column defeated `reading._measurement_among_numbers`, so a headerless
   file published its whole first record as the column names, described
   239 records of 240 and asked nothing: ruling 8 of 2026-09-17 broken
   by one cell, in delimited text and in a workbook alike;
2. a judged placeholder's occurrence total minus its published
   `missing_by_source` count named one cell of a withheld spelling;
3. a form census counted levels that the published levels already
   account for, so subtraction named one held-back label's form --
   ruling 5 bypassed;
4. a column of dates AND timestamps published timestamp-only properties
   against the whole column, so subtracting the date-only cells named
   one exceptional timestamp;
5. `n_numeric` minus `n_negative` minus the plus-sign census named one
   unsigned non-negative cell;
6. a published present-cell population minus the formula count named
   one literal cell;
7. preamble lines and the header line padded a line-ending run over the
   floor, so subtracting them named one record's ending.

Every table and every workbook here is built by neutral code at runtime
(plan D13); the vocabulary is made up on the spot and no cell of anyone's
table appears in this file.
"""

import contextlib
import datetime
import io
import json
import pathlib
import sys

import pytest

from synthtwin import parsing, reading
from tests import workbooks

_FLOOR_ELEVEN = "11"


def _exit_of(argv: "list[str]") -> int:
    """Run one synthtwin command in this process and return its exit code."""
    from synthtwin import cli

    before = sys.argv
    sys.argv = ["synthtwin"] + argv
    try:
        code = cli.main()
    except SystemExit as stop:
        return 0 if stop.code is None else int(stop.code)
    finally:
        sys.argv = before
    return code


def _quiet(argv: "list[str]") -> int:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        return _exit_of(argv)


def _document(path: pathlib.Path) -> dict:
    loaded = json.loads(path.read_text(encoding="utf-8"))
    assert isinstance(loaded, dict)
    return loaded


# -- item 1: one cell of no value reopens publication of a record ------
#
# THE REVIEWER'S INPUT. A headerless file with no furniture above it:
# record 1 is `R001,North Unit,<0.10`, records 2 to 239 are
# `R{i:03},East,{i}.5`, and record 240 is `R240,East,NA`. Floor 11.
#
# MEASURED ON `05e7d89`, the last record's reading changed and nothing
# else, delimited text and workbook alike:
#
#   `240.5`  -> `column_1`..`column_3`, 240 records, the question asked
#   (empty)  -> the same; the old population already dropped the empty
#   `NA`     -> `R001`, `North Unit`, `<0.10` as the names; 239 records
#   `#N/A`   -> the same three names; 239 records
#   `<0.05`  -> the same three names; 239 records
#
# The three failing shapes all now read as records.

_LEAD = "R001,North Unit,<0.10"


def _records_ending_with(reading_text: str) -> str:
    """240 records, no header and no title, the last reading as given."""
    lines = [_LEAD]
    for index in range(2, 240):
        lines = lines + [f"R{index:03},East,{index}.5"]
    return "\n".join(lines + [f"R240,East,{reading_text}"]) + "\n"


def _records_with_eleven(reading_text: str) -> str:
    """The same 240 records with ELEVEN readings written as given.

    WHY THE RED CHECKS BELOW COUNT TO ELEVEN AND NOT TO ONE (the repair
    pass of this landing, round 2's disclosure finding 4). The rule now
    TOLERATES values below that read as neither a spelling of no value
    nor a reading, while they are fewer than `parsing.census_floor` of
    the person's floor -- because one cell somebody typed as `missing`
    used to put a whole record back among the column names. That
    tolerance means a red check that withdraws one seam over a single
    cell is covered by the other, and measures nothing: withdrawing
    `parsing.is_missing_text` over ONE `NA` leaves one unread cell of
    239, the tolerance carries it, and the rule still answers True.

    So each seam is withdrawn over a population the tolerance cannot
    carry: eleven cells at a floor of eleven, which is the line itself.
    Eleven is also the size a real export reaches: the header
    regression that finding 2 closed was measured on a file writing
    `NA` every twenty-ninth row of 320, which is eleven cells. So the
    shape is the reviewer's, at the size that makes the seam load
    bearing. Every twentieth record of 2 to 239 is eleven cells.
    """
    lines = [_LEAD]
    for index in range(2, 240):
        if index % 20 == 0:
            lines = lines + [f"R{index:03},East,{reading_text}"]
        else:
            lines = lines + [f"R{index:03},East,{index}.5"]
    return "\n".join(lines + ["R240,East,240.5"]) + "\n"


def _described(tmp_path: pathlib.Path, body: str) -> dict:
    table = tmp_path / "table.csv"
    table.write_text(body, encoding="utf-8", newline="")
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{tmp_path}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    document = _document(tmp_path / "table-profile.json")
    questions = _document(tmp_path / "table-questions.json")
    from synthtwin import asking

    return {
        "names": [block["name"] for block in document["columns"]],
        "n_rows": document["n_rows"],
        "asked": [
            entry
            for entry in questions["about_your_file"]
            if entry["column"] == asking.FIRST_ROW_SUBJECT
        ],
        "written": (
            (tmp_path / "table-profile.json").read_text(encoding="utf-8")
            + (tmp_path / "table-profile.txt").read_text(encoding="utf-8")
            + (tmp_path / "table-questions.json").read_text(encoding="utf-8")
        ),
    }


@pytest.mark.parametrize("reading_text", ["NA", "#N/A", "<0.05", "", "240.5"])
def test_one_absent_or_censored_reading_never_publishes_a_record(
    tmp_path: pathlib.Path, reading_text: str
) -> None:
    """Item 1, the BLOCKER, reproduced and closed for delimited text.

    Three of these five spellings published the first record as the
    file's three column names on `05e7d89` and described 239 records of
    240; the other two were already right and must stay right.
    """
    read = _described(tmp_path, _records_ending_with(reading_text))
    assert read["names"] == ["column_1", "column_2", "column_3"]
    assert read["n_rows"] == 240
    assert len(read["asked"]) == 1
    seen = read["asked"][0]["what_synthtwin_saw"]
    assert "column 3" in seen and "measurement" in seen
    for text in ("R001", "North Unit", "<0.10"):
        assert text not in read["written"], (
            f"the first record's text {text!r} reached a published document"
        )


def _headerless_workbook(reading_text: str) -> bytes:
    """The same 240 records as a sheet of inline strings."""
    rows = [
        (
            1,
            [
                workbooks.cell("A1", "R001", kind="inlineStr"),
                workbooks.cell("B1", "North Unit", kind="inlineStr"),
                workbooks.cell("C1", "<0.10", kind="inlineStr"),
            ],
        )
    ]
    for index in range(2, 241):
        written = f"{index}.5" if index < 240 else reading_text
        rows = rows + [
            (
                index,
                [
                    workbooks.cell(
                        f"A{index}", f"R{index:03}", kind="inlineStr"
                    ),
                    workbooks.cell(f"B{index}", "East", kind="inlineStr"),
                    workbooks.cell(
                        f"C{index}", written, kind="inlineStr"
                    ),
                ],
            )
        ]
    return workbooks.package(
        [
            (
                "[Content_Types].xml",
                workbooks._content_types(1, False, False, False),
            ),
            ("_rels/.rels", workbooks._root_rels()),
            ("xl/workbook.xml", workbooks._workbook([("Data", "")])),
            (
                "xl/_rels/workbook.xml.rels",
                workbooks._workbook_rels(1, False),
            ),
            ("xl/styles.xml", workbooks._styles()),
            (
                "xl/worksheets/sheet1.xml",
                workbooks.sheet(rows, dimension="A1:C240"),
            ),
        ]
    )


@pytest.mark.parametrize("reading_text", ["NA", "<0.05"])
def test_the_same_workbook_publishes_no_record_either(
    tmp_path: pathlib.Path, reading_text: str
) -> None:
    """Item 1 on the workbook side, through to the twin's header row.

    Measured on `05e7d89`: both spellings gave `R001`, `North Unit` and
    `<0.10` as the three column names over 239 records, and the twin's
    own first sheet row was that person's record verbatim.
    """
    table = tmp_path / "real.xlsx"
    table.write_bytes(_headerless_workbook(reading_text))
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{tmp_path}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    document = _document(tmp_path / "real-profile.json")
    assert [block["name"] for block in document["columns"]] == [
        "column_1", "column_2", "column_3",
    ]
    assert document["n_rows"] == 240
    assert (
        _quiet(
            [
                "generate", f"{tmp_path / 'real-profile.json'}",
                "--out-dir", f"{tmp_path}", "--seed", "4", "--replace",
            ]
        )
        == 0
    )
    import openpyxl

    book = openpyxl.load_workbook(tmp_path / "real-twin.xlsx")
    sheet = book[book.sheetnames[0]]
    header = [
        f"{sheet.cell(row=1, column=place).value}" for place in (1, 2, 3)
    ]
    assert header != ["R001", "North Unit", "<0.10"]
    assert sheet.max_row == 240


def test_the_missing_spellings_filter_is_what_holds_item_one_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 1's first half: the vocabulary of no value.

    With `parsing.is_missing_text` answering False for everything, the
    eleven `NA` cells are back in the evidence population, they reach
    the line the tolerance stops at, the rule declines the column, and
    the file publishes the record again -- 239 records and the three
    real values as the names, which is exactly what `05e7d89` did with
    one such cell.

    The filter in place, the same file reads as records: that is the
    first assertion, and it is what makes the second a withdrawal of
    this seam rather than a statement about the table.
    """
    body = _records_with_eleven("NA")
    assert _described(tmp_path, body)["names"] == [
        "column_1", "column_2", "column_3"
    ]
    monkeypatch.setattr(parsing, "is_missing_text", lambda text: False)
    read = _described(tmp_path, body)
    assert read["names"] == ["R001", "North Unit", "<0.10"]
    assert read["n_rows"] == 239


def test_the_reading_population_is_what_holds_the_censored_cell_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 1's second half: a second censored reading.

    Put `reading._reads_as_a_reading` back to the rule it replaced --
    "this value parses as a number" -- and the eleven RANGE readings
    defeat the column again, republishing the record over 239 rows.
    Eleven and not one, for the reason `_records_with_eleven` gives.

    THE RANGE AND NOT A SECOND CENSORED READING, which is a fact about
    the withdrawal rather than about the rule. Eleven cells written
    `<0.05` are caught by the FOURTH record rule instead of the fifth:
    they wear the same silhouette as `<0.10`, so `_shares_the_shape_
    below` speaks for that column and the file reads as records however
    this seam answers. `2-4` wears another silhouette, so the fifth
    rule is the only one that can speak and the withdrawal measures it.
    """
    body = _records_with_eleven("2-4")
    assert _described(tmp_path, body)["names"] == [
        "column_1", "column_2", "column_3"
    ]
    monkeypatch.setattr(
        reading,
        "_reads_as_a_reading",
        lambda text: parsing.classify_number(text) != parsing.NOT_A_NUMBER,
    )
    read = _described(tmp_path, body)
    assert read["names"] == ["R001", "North Unit", "<0.10"]
    assert read["n_rows"] == 239


def test_a_column_of_readings_with_no_numbers_is_still_no_evidence(
    tmp_path: pathlib.Path
) -> None:
    """The rule stays anchored to a column of NUMBERS.

    Two numbers are required below, as they always were: a column
    written entirely as censored readings offers none for the first
    row's value to stand among, so this rule declines it and the file
    falls through to the furniture rule and to convention. Without that
    floor the widened population would have read any column of
    marked-up text as a column of readings.
    """
    below = ["<0.05"] * 12 + ["NA"] * 3
    assert not reading._measurement_among_numbers("<0.10", below, 11)
    assert reading._measurement_among_numbers("<0.10", ["1.5", "2.5", "NA"], 11)
    assert not reading._measurement_among_numbers("<0.10", ["1.5", "NA"], 11)


# -- the round-trip harness the rest of the items are measured with ---


def _round_trip(
    folder: pathlib.Path,
    body: "str | bytes",
    flags: "list[str]",
    name: str = "table.csv",
    seed: str = "4",
) -> dict:
    """Describe, generate, and validate the twin AND the real file.

    Returns the description and the four exit codes. Every file is
    written with its bytes fixed, so the suite reads the same bytes on
    every platform (the rule `tests/test_description_line_endings.py`
    holds this suite to).
    """
    folder.mkdir(parents=True, exist_ok=True)
    table = folder / name
    if isinstance(body, bytes):
        table.write_bytes(body)
    else:
        table.write_text(body, encoding="utf-8", newline="")
    stem = table.stem
    assert (
        _quiet(
            ["profile", f"{table}", "--out-dir", f"{folder}", "--replace"]
            + flags
        )
        == 0
    )
    described = folder / f"{stem}-profile.json"
    assert (
        _quiet(
            [
                "generate", f"{described}", "--out-dir", f"{folder}",
                "--seed", seed, "--replace",
            ]
        )
        == 0
    )
    twin = folder / f"{stem}-twin{table.suffix}"
    checked_twin = folder / "check-twin"
    checked_twin.mkdir(exist_ok=True)
    checked_real = folder / "check-real"
    checked_real.mkdir(exist_ok=True)
    return {
        "document": _document(described),
        "twin_exit": _quiet(
            [
                "validate", f"{described}", "--twin", f"{twin}",
                "--out-dir", f"{checked_twin}", "--replace",
            ]
        ),
        "real_exit": _quiet(
            [
                "validate", f"{described}", "--twin", f"{table}",
                "--out-dir", f"{checked_real}", "--replace",
            ]
        ),
        "twin": twin,
    }


# -- item 2: a judged candidate's occurrences name a pooled spelling --
#
# THE REVIEWER'S INPUT, at a floor of eleven: a declared measurement of
# 400 cells -- 368 values `100 + i % 100`, twenty `-999`, one `-999.0`
# and eleven `NA`.
#
# MEASURED ON `05e7d89`: `missing_by_source {"-999": 20}` and
# `missing_by_class {"(numeric-sentinel)": 21, "(withheld)": 11}` beside
# a judged `-999` verdict of `n_occurrences` 21 -- so 21 less 20 is one
# cell of a spelling the description never names, twice over. Both files
# validated at nought and the twin wrote twenty `-999` cells and twelve
# blanks.


def _sentinel_column() -> str:
    cells = (
        [f"{100 + place % 100}" for place in range(368)]
        + ["-999"] * 20
        + ["-999.0"]
        + ["NA"] * 11
    )
    rows = ["reading,site"]
    for place in range(len(cells)):
        rows = rows + [f"{cells[place]},s{place % 3}"]
    return "\n".join(rows) + "\n"


def _reading_block(document: dict) -> dict:
    for block in document["columns"]:
        if block["name"] == "reading":
            assert isinstance(block, dict)
            return block
    raise AssertionError("the declared column is not in the description")


def test_a_judged_candidate_leaves_no_spelling_of_one(
    tmp_path: pathlib.Path,
) -> None:
    """Item 2, reproduced and closed.

    The spelling census joins the pool rather than publishing a count
    whose difference from the verdict's own total is one, and both maps
    are raised together, as `_missing_maps` has done since P4-D293.
    """
    read = _round_trip(
        tmp_path / "sentinels",
        _sentinel_column(),
        ["--smallest-group", _FLOOR_ELEVEN, "--measurement", "reading"],
    )
    block = _reading_block(read["document"])
    assert block["missing_by_source"] == {}
    assert block["missing_by_class"]["(withheld)"] == 32
    assert block["n_missing"] == 32
    assert block["n_missing_withheld"] == 32
    verdicts = block["sentinel_verdicts"]
    assert len(verdicts) == 1
    assert verdicts[0]["n_occurrences"] == 21
    assert verdicts[0]["spellings"] == []
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_the_judged_total_is_what_holds_item_two_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 2: withdraw the verdict's own total.

    With `_judged_totals` handing back only the pair the caller already
    held, the pool is not raised and the description publishes
    `{"-999": 20}` beside an `n_occurrences` of 21 again -- the exact
    numbers measured on `05e7d89`.
    """
    from synthtwin import taxonomy

    monkeypatch.setattr(
        taxonomy, "_judged_totals", lambda entries, judged, by_source: []
    )
    folder = tmp_path / "sentinels"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(_sentinel_column(), encoding="utf-8", newline="")
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
                "--measurement", "reading",
            ]
        )
        == 0
    )
    block = _reading_block(_document(folder / "table-profile.json"))
    assert block["missing_by_source"] == {"-999": 20}
    assert block["sentinel_verdicts"][0]["n_occurrences"] == 21
    # ...AND THE LOADER'S HALF REFUSES WHAT THE PRODUCER JUST WROTE,
    # which is the mirror V5 gained with this item: the description
    # does not even reach the generator.
    assert (
        _quiet(
            [
                "generate", f"{folder / 'table-profile.json'}",
                "--out-dir", f"{folder}", "--seed", "4", "--replace",
            ]
        )
        == 1
    )


def _judged_pair(second: int) -> str:
    """375 ordinary readings, twenty `-999`, and `second` of `-999.0`."""
    cells = (
        [f"{100 + place % 100}" for place in range(375)]
        + ["-999"] * 20
        + ["-999.0"] * second
    )
    rows = ["reading,site"]
    for place in range(len(cells)):
        rows = rows + [f"{cells[place]},s{place % 3}"]
    return "\n".join(rows) + "\n"


@pytest.mark.parametrize("second", [2, 5, 10])
def test_a_judged_candidate_leaves_no_group_under_the_floor(
    tmp_path: pathlib.Path, second: int
) -> None:
    """The remainder is nought or reaches the floor (finding 3).

    THE ITEM'S OWN EXPECTATION, which the first repair did not meet.
    It closed a remainder of ONE, because it asked
    `census_names_one_row` at that function's line of two. **Measured**
    at a floor of eleven on this landing's own tree before the repair
    pass: twenty `-999` beside TWO `-999.0` published `missing_by_source
    {"-999": 20, "NA": 11}` against an `n_occurrences` of 22, so 22 less
    20 named two cells wearing a spelling the description never names;
    five `-999.0` named five the same way, and ten named ten. All three
    are groups under a floor of eleven, which is exactly what the floor
    is for.

    Now the pool takes the named spellings until the remainder is
    nought or reaches the line: the spelling census goes empty, the
    verdict names no spelling, and an empty census is asked nothing
    because it leaves the reader nothing to subtract.
    """
    read = _round_trip(
        tmp_path / f"bound{second}",
        _judged_pair(second),
        ["--smallest-group", _FLOOR_ELEVEN, "--measurement", "reading"],
    )
    block = _reading_block(read["document"])
    assert block["missing_by_source"] == {}
    assert block["sentinel_verdicts"][0]["spellings"] == []
    assert block["sentinel_verdicts"][0]["n_occurrences"] == 20 + second
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_a_second_spelling_that_reaches_the_floor_is_published(
    tmp_path: pathlib.Path,
) -> None:
    """The bound stays AT MOST and not EXACTLY, which V5 states.

    Eleven cells of the second spelling reach the line at a floor of
    eleven, so both spellings are named, they cover all 31 of the
    decision's occurrences, and the remainder is nought. Nothing about
    this description moves under finding 3's raised line -- measured
    identically on `05e7d89` and here -- which is what shows the repair
    pools where a reader could subtract and nowhere else.
    """
    read = _round_trip(
        tmp_path / "bound11",
        _judged_pair(11),
        ["--smallest-group", _FLOOR_ELEVEN, "--measurement", "reading"],
    )
    block = _reading_block(read["document"])
    assert block["missing_by_source"] == {"-999": 20, "-999.0": 11}
    assert block["sentinel_verdicts"][0]["n_occurrences"] == 31
    assert block["sentinel_verdicts"][0]["spellings"] == ["-999", "-999.0"]
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_the_default_floor_reads_the_judged_remainder_as_it_always_did(
    tmp_path: pathlib.Path,
) -> None:
    """Finding 3's raised line moves NOTHING at a floor of one.

    `parsing.census_floor(1)` is two, which is the line
    `census_names_one_row` already asked, so the same table describes
    identically at the default floor: both spellings named, the
    remainder nought. Measured on `05e7d89` and here alike.
    """
    read = _round_trip(
        tmp_path / "default",
        _judged_pair(2),
        ["--measurement", "reading"],
    )
    block = _reading_block(read["document"])
    assert block["missing_by_source"] == {"-999": 20, "-999.0": 2}
    assert block["sentinel_verdicts"][0]["spellings"] == ["-999", "-999.0"]
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


# -- item 3: a form census names a held-back label's form -------------
#
# THE REVIEWER'S INPUT, at a floor of eleven and seed 4: one column of
# `ABC-100` x100, `ABC-200` x100, `ABC-300` x1, `QQ-400` x5, `RR-500` x6.
#
# MEASURED ON `05e7d89`: both common levels published
# `shape_form_cells` 100 beside `shape_forms {"@@@-%%%": 201,
# "@@-%%%": 11}`, so 201 less 200 is one held-back level of one row with
# its form on the page. `n_missing` was nought, the twin wrote one
# `AAA-000`, and both files validated with 47 held and nothing missed.


def _form_column() -> str:
    cells = (
        ["ABC-100"] * 100
        + ["ABC-200"] * 100
        + ["ABC-300"]
        + ["QQ-400"] * 5
        + ["RR-500"] * 6
    )
    return "\n".join(["code"] + cells) + "\n"


def test_a_form_census_names_no_held_back_label(
    tmp_path: pathlib.Path,
) -> None:
    """Item 3, reproduced and closed by ruling 5 (2026-09-17, item 5).

    The singleton level's cell is counted as MISSING, so the census
    counts 200 where the published levels cover 200 and the difference
    is nought. The second key is left exactly as it was: its eleven
    cells belong to two held-back levels, no published level wears that
    form, and the pool is doing what the ruling asks of it.
    """
    read = _round_trip(
        tmp_path / "forms",
        _form_column(),
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    block = read["document"]["columns"][0]
    assert block["shape_forms"] == {"@@@-%%%": 200, "@@-%%%": 11}
    covered = 0
    for level in block["levels"]:
        covered = covered + level["shape_form_cells"]
    assert covered == 200
    assert block["n_missing"] == 1
    assert block["suppressed_levels"] == 2
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_the_form_pass_is_what_holds_item_three_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 3: withdraw the form reading of ruling 5.

    With `_levels_read_by_form` counting nothing out, the census
    publishes 201 against 200 published cells again and `n_missing`
    goes back to nought -- `05e7d89` exactly. The loader's own half
    then refuses the description, so it never reaches a twin.
    """
    from synthtwin import taxonomy

    monkeypatch.setattr(
        taxonomy, "_levels_read_by_form", lambda cells, role, rare: ()
    )
    folder = tmp_path / "forms"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(_form_column(), encoding="utf-8", newline="")
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    block = _document(folder / "table-profile.json")["columns"][0]
    assert block["shape_forms"]["@@@-%%%"] == 201
    assert block["n_missing"] == 0
    assert (
        _quiet(
            [
                "generate", f"{folder / 'table-profile.json'}",
                "--out-dir", f"{folder}", "--seed", "4", "--replace",
            ]
        )
        == 1
    )


def test_a_form_census_leaving_two_cells_over_is_unmoved(
    tmp_path: pathlib.Path,
) -> None:
    """The rule fires at ONE cell and at no other number.

    The reviewer's column with a SECOND singleton level of the same
    form leaves two cells over the published levels, which is a group
    and not a row, so nothing is counted out, `n_missing` stays nought
    and the census publishes 202. A rule that counted out every
    held-back level would have emptied this column for nothing.
    """
    cells = (
        ["ABC-100"] * 100
        + ["ABC-200"] * 100
        + ["ABC-300"]
        + ["ABC-400"]
        + ["QQ-400"] * 5
        + ["RR-500"] * 6
    )
    read = _round_trip(
        tmp_path / "two",
        "\n".join(["code"] + cells) + "\n",
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    block = read["document"]["columns"][0]
    assert block["shape_forms"] == {"@@@-%%%": 202, "@@-%%%": 11}
    assert block["n_missing"] == 0
    assert block["suppressed_levels"] == 4
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


# -- item 4: a joint column's timestamp-only residuals ----------------
#
# THE REVIEWER'S TWO INPUTS, at a floor of eleven, each 400 cells:
#
# * 100 ISO dates, 299 timestamps at noon, one at midnight -- MEASURED
#   ON `05e7d89`: `resolution_mix {"iso-date": 100, "iso-datetime":
#   300}` beside `n_at_midnight` 101, and 101 less 100 is one person's
#   time of day;
# * 100 ISO dates, 299 noon timestamps ending `Z`, one unzoned noon
#   timestamp -- MEASURED: the same mix beside `utc_offsets
#   {"(none)": 101, "Z": 299}`, and 101 less 100 is that one cell.
#
# Both descriptions loaded and both source files validated with nothing
# missed.

_FIRST_DAY = datetime.date(2021, 1, 1)


def _day(place: int) -> str:
    return (_FIRST_DAY + datetime.timedelta(days=place)).isoformat()


def _joint_column(odd_one: str) -> str:
    cells: "list[str]" = []
    for place in range(100):
        cells = cells + [_day(place)]
    for place in range(299):
        cells = cells + [f"{_day(place)}T12:00:00"]
    return "\n".join(["moment"] + cells + [odd_one]) + "\n"


def _zoned_column(odd_one: str) -> str:
    cells: "list[str]" = []
    for place in range(100):
        cells = cells + [_day(place)]
    for place in range(299):
        cells = cells + [f"{_day(place)}T12:00:00Z"]
    return "\n".join(["moment"] + cells + [odd_one]) + "\n"


def test_a_joint_column_names_no_lone_midnight_timestamp(
    tmp_path: pathlib.Path,
) -> None:
    """Item 4's first half, reproduced and closed.

    The whole dates stand at midnight by definition and the mix counts
    them exactly, so the count is held to the floor on both sides of
    the TIMESTAMPS' own residual -- and where it cannot be, the field
    goes to its one silence, `null`.
    """
    read = _round_trip(
        tmp_path / "midnight",
        _joint_column(f"{_day(5)}T00:00:00"),
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    block = read["document"]["columns"][0]
    assert block["resolution_mix"] == {"iso-date": 100, "iso-datetime": 300}
    assert block["n_at_midnight"] is None
    assert block["all_at_midnight"] is False
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_a_joint_column_names_no_lone_unzoned_timestamp(
    tmp_path: pathlib.Path,
) -> None:
    """Item 4's second half, reproduced and closed by ruling 6.

    The rare offset of the TIMESTAMPS is counted into the commonest one
    they wrote, so the one unzoned cell is read at `Z` and `(none)`
    counts the whole dates and nothing else: 100 against a mix of 100.
    """
    read = _round_trip(
        tmp_path / "offsets",
        _zoned_column(f"{_day(5)}T12:00:00"),
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    block = read["document"]["columns"][0]
    assert block["resolution_mix"] == {"iso-date": 100, "iso-datetime": 300}
    assert block["utc_offsets"] == {"(none)": 100, "Z": 300}
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_the_midnight_residual_is_what_holds_item_four_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 4's first half.

    With `_joint_midnight_nameable` answering True for everything the
    count comes back as 101 beside 100 date-only cells -- the number
    measured on `05e7d89` -- and the loader's own half then refuses it.
    """
    from synthtwin import taxonomy

    monkeypatch.setattr(
        taxonomy,
        "_joint_midnight_nameable",
        lambda format_name, sources, floor: True,
    )
    folder = tmp_path / "midnight"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(
        _joint_column(f"{_day(5)}T00:00:00"), encoding="utf-8", newline=""
    )
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    block = _document(folder / "table-profile.json")["columns"][0]
    assert block["n_at_midnight"] == 101
    assert (
        _quiet(
            [
                "generate", f"{folder / 'table-profile.json'}",
                "--out-dir", f"{folder}", "--seed", "4", "--replace",
            ]
        )
        == 1
    )


def test_the_timestamp_population_is_what_holds_the_offset_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 4's second half.

    With `_date_only_places` finding none, the offset census is decided
    over the whole column again, `(none)` clears the floor on the whole
    dates' backs, and the map is `{"(none)": 101, "Z": 299}` -- the
    numbers measured on `05e7d89`. The loader then refuses it.
    """
    from synthtwin import taxonomy

    monkeypatch.setattr(
        taxonomy,
        "_date_only_places",
        lambda format_name, pairs, sources: {},
    )
    folder = tmp_path / "offsets"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(
        _zoned_column(f"{_day(5)}T12:00:00"), encoding="utf-8", newline=""
    )
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    block = _document(folder / "table-profile.json")["columns"][0]
    assert block["utc_offsets"] == {"(none)": 101, "Z": 299}
    assert (
        _quiet(
            [
                "generate", f"{folder / 'table-profile.json'}",
                "--out-dir", f"{folder}", "--seed", "4", "--replace",
            ]
        )
        == 1
    )


def test_a_joint_column_of_ordinary_size_still_counts_that_hour(
    tmp_path: pathlib.Path,
) -> None:
    """A residual that is a GROUP on both sides still publishes.

    100 whole dates, 150 timestamps at midnight and 150 at noon leave
    150 and 150, so the count stands at 250 and nothing is withheld.
    """
    cells: "list[str]" = []
    for place in range(100):
        cells = cells + [_day(place)]
    for place in range(150):
        cells = cells + [f"{_day(place)}T00:00:00"]
    for place in range(150):
        cells = cells + [f"{_day(place)}T12:00:00"]
    read = _round_trip(
        tmp_path / "even",
        "\n".join(["moment"] + cells) + "\n",
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    block = read["document"]["columns"][0]
    assert block["resolution_mix"] == {"iso-date": 100, "iso-datetime": 300}
    assert block["n_at_midnight"] == 250
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


# -- item 5: the plus census and the population a reader derives ------
#
# THE REVIEWER'S INPUT, at a floor of eleven: a declared measurement of
# 500 cells -- 400 written `+100.5` upward, 99 written `-100.5` downward
# and ONE written `250.5`.
#
# MEASURED ON `05e7d89`: `n_numeric` 500, `n_negative` 99 and
# `decimal_plus {"+": 400}`, so 500 less 99 less 400 is one unsigned
# non-negative cell. The description loaded and the source file
# validated with 69 held and nothing missed.


def _signed_column() -> str:
    cells: "list[str]" = []
    for place in range(400):
        cells = cells + [f"+{100 + place}.5"]
    for place in range(99):
        cells = cells + [f"-{100 + place}.5"]
    return "\n".join(["reading"] + cells + ["250.5"]) + "\n"


def _profiled(
    tmp_path: pathlib.Path, body: str, flags: "list[str]"
) -> dict:
    """Describe one file and hand back its first column's block."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    table = tmp_path / "table.csv"
    table.write_text(body, encoding="utf-8", newline="")
    assert (
        _quiet(
            ["profile", f"{table}", "--out-dir", f"{tmp_path}", "--replace"]
            + flags
        )
        == 0
    )
    block = _document(tmp_path / "table-profile.json")["columns"][0]
    assert isinstance(block, dict)
    return block


def test_the_plus_census_leaves_no_unsigned_cell_of_one(
    tmp_path: pathlib.Path,
) -> None:
    """Item 5, reproduced and closed by ruling 6 (2026-09-17, item 6).

    The one unsigned non-negative cell is a spelling below the line, so
    it is counted into the commonest spelling the column wrote -- the
    plus -- and the census says 401. The subtraction then comes to
    nought, and 400 cells keep the plus the twin writes them with.
    """
    block = _profiled(
        tmp_path / "signed",
        _signed_column(),
        ["--smallest-group", _FLOOR_ELEVEN, "--measurement", "reading"],
    )
    assert block["n_numeric"] == 500
    assert block["n_negative"] == 99
    assert block["decimal_plus"] == {"+": 401}
    assert block["n_numeric"] - block["n_negative"] - 401 == 0


def test_the_derived_population_is_what_holds_item_five_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 5: withdraw the population a reader derives.

    With `_signable_decimals` answering nought, neither ruling 6 nor
    the second reading of the disclosure rule is asked, and the census
    publishes 400 against 500 decimals and 99 negatives -- `05e7d89`
    exactly, one cell recoverable by subtraction.
    """
    from synthtwin import taxonomy

    monkeypatch.setattr(
        taxonomy, "_signable_decimals", lambda decimals, negatives: 0
    )
    block = _profiled(
        tmp_path / "signed",
        _signed_column(),
        ["--smallest-group", _FLOOR_ELEVEN, "--measurement", "reading"],
    )
    assert block["decimal_plus"] == {"+": 400}
    assert block["n_numeric"] - block["n_negative"] - 400 == 1


def test_a_column_of_both_kinds_of_number_is_unmoved(
    tmp_path: pathlib.Path,
) -> None:
    """The derived population is read only where it is consistent.

    `n_negative` counts the negative WHOLE numbers too, so on a column
    written half `+12` and half `+3.25` the difference falls below the
    signed count and says nothing. It is not read there, and the census
    is the exact count of signed decimals it always was.
    """
    cells: "list[str]" = []
    for place in range(450):
        cells = cells + [f"{place - 225:+d}"]
    for place in range(450):
        cells = cells + [f"{(place - 225) + 0.25:+.2f}"]
    block = _profiled(
        tmp_path / "both",
        "\n".join(["change"] + cells) + "\n",
        ["--smallest-group", _FLOOR_ELEVEN, "--measurement", "change"],
    )
    signed = 0
    for cell in cells:
        if cell[:1] == "+" and "." in cell:
            signed = signed + 1
    assert block["decimal_plus"] == {"+": signed}


# -- item 6: a formula count names a lone literal cell ----------------
#
# THE REVIEWER'S INPUT, at a floor of eleven: a 400-record workbook
# whose measurement column holds twenty formulas `20+1` through `20+20`
# with cached results 21 to 40, ONE literal `41`, and 379 physically
# absent cells; a second column is populated throughout.
#
# MEASURED ON `05e7d89`: `cell_classes {"number": 21, "absent": 379}`
# with every other class nought, beside `formulas` 20 -- and 21 less 20
# is the one cell somebody typed by hand. Both ends of 20 against 400
# rows clear eleven, which is why the row count alone let it through.


def _formula_book(literal: int = 1, n_rows: int = 400) -> bytes:
    rows = [
        (
            1,
            [
                workbooks.cell("A1", "reading", kind="inlineStr"),
                workbooks.cell("B1", "site", kind="inlineStr"),
            ],
        )
    ]
    for index in range(2, n_rows + 2):
        record = index - 1
        cells: "list[str]" = []
        if record <= 20:
            cells = cells + [
                workbooks.cell(
                    f"A{index}", f"{20 + record}", formula=f"20+{record}"
                )
            ]
        elif record <= 20 + literal:
            cells = cells + [
                workbooks.cell(f"A{index}", f"{20 + record}")
            ]
        cells = cells + [
            workbooks.cell(f"B{index}", f"s{record % 4}", kind="inlineStr")
        ]
        rows = rows + [(index, cells)]
    return workbooks.package(
        [
            (
                "[Content_Types].xml",
                workbooks._content_types(1, False, False, False),
            ),
            ("_rels/.rels", workbooks._root_rels()),
            ("xl/workbook.xml", workbooks._workbook([("Data", "")])),
            (
                "xl/_rels/workbook.xml.rels",
                workbooks._workbook_rels(1, False),
            ),
            ("xl/styles.xml", workbooks._styles()),
            (
                "xl/worksheets/sheet1.xml",
                workbooks.sheet(rows, dimension=f"A1:B{n_rows + 1}"),
            ),
        ]
    )


def _sheet_column(document: dict, place: int) -> dict:
    block = document["source"]["workbook"]["columns"][place]
    assert isinstance(block, dict)
    return block


def test_a_formula_count_names_no_literal_cell(
    tmp_path: pathlib.Path,
) -> None:
    """Item 6, reproduced and closed.

    A formula cell is a PRESENT cell, so the count is held to the line
    inside the present cells the census publishes as well as inside the
    rows -- and twenty of twenty-one leaves one, so the count is
    withheld.
    """
    read = _round_trip(
        tmp_path / "formulas",
        _formula_book(),
        ["--smallest-group", _FLOOR_ELEVEN],
        name="real.xlsx",
    )
    block = _sheet_column(read["document"], 0)
    assert block["cell_classes"]["number"] == 21
    assert block["cell_classes"]["absent"] == 379
    assert block["formulas"] is None
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_the_present_cells_are_what_hold_item_six_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 6: withdraw the present-cell population.

    With `workbook._present_cells` answering None the count is asked
    only of the 400 rows again, both ends of 20 clear eleven, and the
    description publishes `formulas` 20 beside 21 number cells --
    `05e7d89` exactly. The loader's own half then refuses it.
    """
    from synthtwin import workbook

    monkeypatch.setattr(workbook, "_present_cells", lambda census: None)
    folder = tmp_path / "formulas"
    folder.mkdir(parents=True)
    table = folder / "real.xlsx"
    table.write_bytes(_formula_book())
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    block = _sheet_column(_document(folder / "real-profile.json"), 0)
    assert block["formulas"] == 20
    assert block["cell_classes"]["number"] == 21
    assert (
        _quiet(
            [
                "generate", f"{folder / 'real-profile.json'}",
                "--out-dir", f"{folder}", "--seed", "4", "--replace",
            ]
        )
        == 1
    )


def test_a_formula_count_with_room_beside_it_still_publishes(
    tmp_path: pathlib.Path,
) -> None:
    """A count with a group on both sides inside the present cells stands.

    The same workbook with fifteen literal cells beside the twenty
    formulas leaves 35 present cells and fifteen over, which is a
    group, so the count is published exactly as before.
    """
    read = _round_trip(
        tmp_path / "room",
        _formula_book(literal=15),
        ["--smallest-group", _FLOOR_ELEVEN],
        name="real.xlsx",
    )
    block = _sheet_column(read["document"], 0)
    assert block["cell_classes"]["number"] == 35
    assert block["formulas"] == 20
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


# -- item 7: file furniture disguises a singleton record ending -------
#
# THE REVIEWER'S INPUT, at a floor of eleven: ten title lines, the
# header `record,reading`, record 1 written with a bare newline, and
# records 2 to 120 written with a carriage return and a newline.
#
# MEASURED ON `05e7d89`: a preamble of ten lines beside `line_endings
# [{"lf": 12}, {"crlf": 119}]` -- and 12 less 10 less 1 is the first
# record, the only one of the 120 written that way.


def _furnished_file(titles: int = 10) -> str:
    lines: "list[str]" = []
    for place in range(titles):
        lines = lines + [f"# extract note {place}\n"]
    lines = lines + ["record,reading\n", "R001,1\n"]
    for place in range(2, 121):
        lines = lines + [f"R{place:03},{place}\r\n"]
    return "".join(lines)


def _form_of(document: dict) -> dict:
    form = document["source"]["dialect"]
    assert isinstance(form, dict)
    return form


def test_file_furniture_never_carries_a_lone_record_ending(
    tmp_path: pathlib.Path,
) -> None:
    """Item 7, reproduced and closed.

    The lines above the table come off the front before the run rule is
    asked, so the file collapses to one run of the commonest ending --
    which is what P4-D290 already did for every other rare ending.
    """
    read = _round_trip(
        tmp_path / "furniture",
        _furnished_file(),
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    form = _form_of(read["document"])
    assert form["preamble"] == [{"kind": "comment", "lines": 10, "mark": "# "}]
    assert form["line_endings"] == [{"ending": "crlf", "lines": 131}]
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_the_furniture_subtraction_is_what_holds_item_seven_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for item 7: withdraw the subtraction.

    With `record_endings` handing the runs back untouched, the ten
    title lines and the header carry the bare newline over the line
    again and the description publishes `[{lf: 12}, {crlf: 119}]` --
    `05e7d89` exactly. The loader's own half then refuses it.
    """
    from synthtwin import dialect

    monkeypatch.setattr(
        dialect, "record_endings", lambda runs, above: list(runs)
    )
    folder = tmp_path / "furniture"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(_furnished_file(), encoding="utf-8", newline="")
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}",
                "--replace", "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    form = _form_of(_document(folder / "table-profile.json"))
    assert form["line_endings"] == [
        {"ending": "lf", "lines": 12}, {"ending": "crlf", "lines": 119},
    ]
    # ...AND THE LOADER'S HALF REFUSES IT once the rule is back in
    # place. It is put back first because `dialect.endings_broken` asks
    # the very function this test withdrew, which is what makes the two
    # halves one rule rather than two.
    monkeypatch.undo()
    assert (
        _quiet(
            [
                "generate", f"{folder / 'table-profile.json'}",
                "--out-dir", f"{folder}", "--seed", "4", "--replace",
            ]
        )
        == 1
    )


def test_a_file_whose_records_change_ending_in_groups_is_unmoved(
    tmp_path: pathlib.Path,
) -> None:
    """Runs that are groups on both sides of the furniture still stand.

    Ten title lines, a header, then sixty records with a bare newline
    and sixty with a carriage return and a newline: the records' own
    runs are sixty and sixty, so nothing collapses and the file's form
    is published as it stands.
    """
    lines: "list[str]" = []
    for place in range(10):
        lines = lines + [f"# extract note {place}\n"]
    lines = lines + ["record,reading\n"]
    for place in range(1, 61):
        lines = lines + [f"R{place:03},{place}\n"]
    for place in range(61, 121):
        lines = lines + [f"R{place:03},{place}\r\n"]
    read = _round_trip(
        tmp_path / "groups",
        "".join(lines),
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    form = _form_of(read["document"])
    assert form["line_endings"] == [
        {"ending": "lf", "lines": 71}, {"ending": "crlf", "lines": 60},
    ]
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


# =====================================================================
# THE REPAIR PASS OF THIS LANDING: the skeptic's five results.
#
# The seven items above were verified against `05e7d89` and against the
# landing's own commit by a skeptic reading for what the REPAIRS cost.
# Five came back and each is closed here, from the skeptic's
# own reproduction, with its own red check.
#
#   1 BLOCKER -- item 4's midnight repair made the tool generate a twin
#     that then failed its own description at the DEFAULT floor;
#   2 MAJOR   -- item 1's widening threw away the real column names of
#     ordinary headed exports whose header carries a figure;
#   3 MINOR   -- item 2 closed a remainder of one and left two to ten;
#   4 MINOR   -- one cell somebody typed as `missing` still published a
#     whole record as the column names;
#   5 MINOR   -- no KPI entry moved when any repair was withdrawn.
#
# Finding 3 is closed beside item 2 above; finding 5 is closed in
# `tests/kpi/ledger.json` and checked by `tests/test_kpi_ledger.py`.
# =====================================================================


def _joint_iso_column(at_midnight: int, zulu: int, unzoned: int) -> str:
    """100 whole ISO dates over timestamps, as the reviewer wrote them."""
    lines = ["when"]
    for place in range(100):
        lines = lines + [f"2024-01-{(place % 28) + 1:02}"]
    day = 0
    for place in range(at_midnight):
        day = day + 1
        lines = lines + [f"2024-02-{(place % 28) + 1:02}T00:00:00"]
    for place in range(zulu):
        lines = lines + [f"2024-03-{(place % 28) + 1:02}T12:00:00Z"]
    for place in range(unzoned):
        lines = lines + [f"2024-03-{(place % 28) + 1:02}T12:00:00"]
    return "\n".join(lines) + "\n"


def _noon_and_one_midnight() -> str:
    """The reviewer's item-4(a) input: 100 dates, 299 noon, one midnight."""
    lines = ["when"]
    for place in range(100):
        lines = lines + [f"2024-01-{(place % 28) + 1:02}"]
    for place in range(299):
        lines = lines + [f"2024-02-{(place % 28) + 1:02}T12:00:00"]
    return "\n".join(lines + ["2024-02-04T00:00:00"]) + "\n"


@pytest.mark.parametrize("floor", [[], ["--smallest-group", "2"],
                                   ["--smallest-group", "5"],
                                   ["--smallest-group", _FLOOR_ELEVEN]])
@pytest.mark.parametrize("seed", ["0", "4", "19"])
def test_a_joint_column_generates_a_twin_that_passes_its_own_description(
    tmp_path: pathlib.Path, floor: "list[str]", seed: str
) -> None:
    """FINDING 1, the blocker of the repair pass.

    THE SKEPTIC'S REPRODUCTION. Item 4(a)'s own input at
    `--smallest-group 1`, which is the DEFAULT: 100 whole ISO dates,
    299 timestamps at noon and one at midnight. Item 4 withholds
    `n_at_midnight` there, correctly -- 101 less the hundred whole
    dates is one person's time of day. **Measured** on the landing's
    own commit before this repair: `synthtwin generate` wrote a twin
    and `synthtwin validate` then exited 3 on `midnight.withheld
    [datetime.n_at_midnight]: MISSED`, at floors one, two and five and
    on every seed tried. `05e7d89` exited 0 at every floor. Floors
    eleven and twenty-five passed only by luck, the twin landing inside
    the band.

    The round's own requirement is that a round trip validate the twin
    AND the real table at exit 0. It does now, at four floors and three
    seeds, and the count stays withheld while it does.
    """
    read = _round_trip(
        tmp_path / f"joint{seed}{len(floor)}",
        _noon_and_one_midnight(),
        floor,
        seed=seed,
    )
    block = read["document"]["columns"][0]
    assert block["n_at_midnight"] is None
    assert block["all_at_midnight"] is False
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_a_single_format_column_still_owes_its_withheld_count(
    tmp_path: pathlib.Path,
) -> None:
    """The blocker's repair reaches JOINT columns and nothing else.

    A column of moments ALONE -- no whole dates in it -- whose midnight
    count is withheld for its size still owes that count on one side,
    which is what P4-D191 built and what the generator's rank shift is
    for. Here 396 timestamps stand across the working day and four at
    midnight, so at a floor of eleven the count is withheld: fewer than
    the line stood there. `resolution_mix` holds ONE key, the guard
    does not fire, and the obligation is still owed and still met.
    """
    lines = ["when"]
    for place in range(396):
        lines = lines + [f"2024-02-{(place % 28) + 1:02}T{9 + place % 8:02}:30:00"]
    for place in range(4):
        lines = lines + [f"2024-03-{place + 1:02}T00:00:00"]
    read = _round_trip(
        tmp_path / "alone",
        "\n".join(lines) + "\n",
        ["--smallest-group", _FLOOR_ELEVEN],
    )
    block = read["document"]["columns"][0]
    assert block["n_at_midnight"] is None
    assert len(block["resolution_mix"]) == 1
    from synthtwin import contract

    loaded = contract.load_profile(
        f"{tmp_path / 'alone' / 'table-profile.json'}"
    )
    facts = loaded.columns[0].facts
    assert facts is not None
    assert contract.midnight_withheld_for_its_size(facts) is True
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


def test_the_joint_guard_is_what_holds_finding_one_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for finding 1: withdraw the joint reading of a silence.

    Put `contract.midnight_withheld_for_its_size` back to the reading
    that cannot tell item 4's new silence from the old one -- the same
    answer it gave before the guard -- and the generator owes an
    obligation no rank shift can meet on a joint column, so the twin
    fails its own description at exit 3 on
    `midnight.withheld [datetime.n_at_midnight]`.
    """
    from synthtwin import contract

    original = contract.midnight_withheld_for_its_size

    def _without_the_guard(facts: object) -> bool:
        mix = getattr(facts, "resolution_mix")
        try:
            object.__setattr__(facts, "resolution_mix", {"iso-datetime": 1})
            return original(facts)  # type: ignore[arg-type]
        finally:
            object.__setattr__(facts, "resolution_mix", mix)

    monkeypatch.setattr(
        contract, "midnight_withheld_for_its_size", _without_the_guard
    )
    folder = tmp_path / "withdrawn"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(_noon_and_one_midnight(), encoding="utf-8", newline="")
    assert (
        _quiet(
            ["profile", f"{table}", "--out-dir", f"{folder}", "--replace"]
        )
        == 0
    )
    described = folder / "table-profile.json"
    assert (
        _quiet(
            [
                "generate", f"{described}", "--out-dir", f"{folder}",
                "--seed", "4", "--replace",
            ]
        )
        == 0
    )
    checked = folder / "checked"
    checked.mkdir()
    assert (
        _quiet(
            [
                "validate", f"{described}",
                "--twin", f"{folder / 'table-twin.csv'}",
                "--out-dir", f"{checked}", "--replace",
            ]
        )
        == 3
    )
    report = (checked / "table-twin-quality.txt").read_text(encoding="utf-8")
    assert "n_at_midnight" in report


def test_a_joint_column_still_names_no_lone_timestamp_at_the_default_floor(
    tmp_path: pathlib.Path,
) -> None:
    """Finding 1's repair does not reopen item 4, at any floor.

    The leak item 4 closed is the arithmetic `n_at_midnight` less the
    exact `resolution_mix` count of whole dates. It stays shut at the
    default floor: the count is absent, so there is nothing to subtract
    from, and the offset census counts the one unzoned timestamp into
    the commonest offset the timestamps wrote.
    """
    read = _round_trip(
        tmp_path / "shut",
        _joint_iso_column(0, 299, 1),
        [],
    )
    block = read["document"]["columns"][0]
    dates = block["resolution_mix"]["iso-date"]
    assert dates == 100
    assert block["utc_offsets"] == {"(none)": 100, "Z": 300}
    assert (read["twin_exit"], read["real_exit"]) == (0, 0)


# -- finding 2: an ordinary headed export loses its column names ------
#
# THE SKEPTIC'S REPRODUCTION. A perfectly ordinary headed CSV: the
# header `site,2024 total`, 320 records `North Unit,437`, with `NA`
# every twenty-ninth row -- eleven cells of 320.
#
# MEASURED: `05e7d89` published `site` and `2024 total` over 320
# records; the landing's own commit published `column_1` and `column_2`
# over 321, the header row described and generated as a 321st record.
# Five of eleven realistic header names flipped that way.

_SITES = ("North Unit", "South Unit", "East Unit", "West Unit")


def _headed_export(header: str, absent: str) -> str:
    """320 records under a header, with `absent` every 29th row."""
    lines = [header]
    for place in range(1, 321):
        cell = absent if place % 29 == 0 else f"{10 + (place * 7) % 890}"
        lines = lines + [f"{_SITES[place % 4]},{cell}"]
    return "\n".join(lines) + "\n"


@pytest.mark.parametrize(
    "header",
    [
        "site,2024 total",
        "site,1st reading",
        "site,100m time",
        "site,3-month change",
        "site,>65 count",
        "site,5 mg dose",
        "site,2-4 week",
    ],
)
@pytest.mark.parametrize("absent", ["NA", "<10", "437"])
@pytest.mark.parametrize("floor", [[], ["--smallest-group", _FLOOR_ELEVEN]])
def test_a_headed_export_keeps_the_names_its_file_wrote(
    tmp_path: pathlib.Path, header: str, absent: str, floor: "list[str]"
) -> None:
    """FINDING 2: the header is a column name, not somebody's reading.

    Each of these headers opens on a figure or a mark and carries one,
    so `_holds_a_figure_as_a_value` accepts it and the fifth record
    rule used to read the header row as a record. The guard asks
    whether the value carries a WORD, which every column name here does
    and no reading does.

    Two of the owner's mandatory goals ride on this: the twin's columns
    must be the real table's columns, or code written on the twin does
    not run unchanged on the real table, and the twin must hold 320
    records, or counts taken on it are wrong.
    """
    folder = tmp_path / f"headed{len(header)}{absent}{len(floor)}"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(
        _headed_export(header, absent), encoding="utf-8", newline=""
    )
    assert (
        _quiet(
            ["profile", f"{table}", "--out-dir", f"{folder}", "--replace"]
            + floor
        )
        == 0
    )
    document = _document(folder / "table-profile.json")
    names = [column["name"] for column in document["columns"]]
    assert names == header.split(",")
    assert document["n_rows"] == 320


def test_the_word_guard_is_what_holds_finding_two_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for finding 2: withdraw the word test.

    With `reading._carries_a_word` answering False the header `2024
    total` is read as a record again: the file's own column names are
    thrown away for `column_1` and `column_2`, and its 320 records
    become 321. Those are the numbers measured on this landing's commit
    before the repair pass.
    """
    monkeypatch.setattr(reading, "_carries_a_word", lambda text: False)
    folder = tmp_path / "withdrawn"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(
        _headed_export("site,2024 total", "NA"), encoding="utf-8", newline=""
    )
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}", "--replace",
                "--smallest-group", _FLOOR_ELEVEN,
            ]
        )
        == 0
    )
    document = _document(folder / "table-profile.json")
    names = [column["name"] for column in document["columns"]]
    assert names == ["column_1", "column_2"]
    assert document["n_rows"] == 321


def test_a_word_is_three_letters_and_a_unit_is_two() -> None:
    """Where the line stands, stated on the values it was drawn for.

    The readings the fifth record rule exists for carry a mark, a range
    or a unit; a column name carries a word. Two letters is a unit and
    three is a word, which is the whole of the rule.
    """
    for reading_text in ["<0.10", "2-4", "5 mg", "250.5", "12.5 kg", "-3"]:
        assert not reading._carries_a_word(reading_text)
    for name in [
        "2024 total", "1st reading", "100m time", "3-month change",
        "><65 count", "5 mg dose", "2-4 week",
    ]:
        assert reading._carries_a_word(name)


# -- finding 4: one unrecognised word publishes a whole record --------
#
# THE SKEPTIC'S REPRODUCTION. The reviewer's own item-1 table with the
# last reading written `missing` instead of `NA`.
#
# MEASURED on `05e7d89` AND on this landing's commit alike: names
# `R001`, `North Unit`, `<0.10`; 239 records; no first-row question;
# all three of that person's values in the description, the printed
# page and the questions file. `parsing.is_missing_text('missing')` is
# False -- so are `MISSING`, `unknown`, `N.A.` and `nil` -- so item 1's
# vocabulary never reached it.


@pytest.mark.parametrize(
    "written",
    ["missing", "MISSING", "Missing", "unknown", "N.A.", "nil",
     "not recorded"],
)
def test_one_word_nobody_declared_never_publishes_a_record(
    tmp_path: pathlib.Path, written: str
) -> None:
    """FINDING 4: the blocker's own class, for text nobody declared.

    The rule counts the values below that read as NEITHER a spelling of
    no value nor a reading, and survives them while they are fewer than
    `parsing.census_floor` of the floor. One such cell in 239 is far
    under the line of eleven, so the record is read as a record: three
    placeholder names, 240 records, the first-row question asked, and
    no value of that record anywhere in what synthtwin publishes.
    """
    read = _described(tmp_path, _records_ending_with(written))
    assert read["names"] == ["column_1", "column_2", "column_3"]
    assert read["n_rows"] == 240
    assert read["asked"]
    for secret in ("R001", "North Unit", "<0.10"):
        assert secret not in read["written"]


def test_a_column_mostly_of_words_is_still_no_column_of_readings(
    tmp_path: pathlib.Path
) -> None:
    """The tolerance is a MINORITY and the floor says how small.

    Eleven cells of an undeclared word at a floor of eleven reach the
    line, so the rule declines the column: a column a reader would call
    text is not read as a column of readings just because two of its
    cells parse as numbers. The file then falls to the furniture rule
    and to convention, as it did before item 1.
    """
    below = ["1.5", "2.5"] + ["not recorded"] * 11
    assert not reading._measurement_among_numbers("<0.10", below, 11)
    assert reading._measurement_among_numbers("<0.10", below[:12], 11)


def test_the_tolerance_is_what_holds_finding_four_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for finding 4: withdraw the minority.

    Hold the tolerance to nought -- which is the rule as item 1 wrote
    it, answering False at the first value that reads as neither -- and
    the single `missing` cell publishes the whole record as the column
    names over 239 records again, which is what `05e7d89` and this
    landing's commit both did.
    """
    monkeypatch.setattr(parsing, "census_floor", lambda floor: 1)
    read = _described(tmp_path, _records_ending_with("missing"))
    assert read["names"] == ["R001", "North Unit", "<0.10"]
    assert read["n_rows"] == 239


def test_the_raised_line_is_what_holds_finding_three_shut(
    tmp_path: pathlib.Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """RED CHECK for finding 3: put the judged reading back to two.

    With `parsing.census_names_one_row` deaf to the line its caller
    names, the judged pair is read at the line of two again: twenty
    `-999` beside two `-999.0` republish `missing_by_source` with
    `-999` in it beside an `n_occurrences` of 22, and 22 less 20 names
    two cells of a spelling the description never names. The loader's
    own V5 then refuses the description the producer just wrote, which
    is what makes the two one rule.
    """
    original = parsing.census_names_one_row
    monkeypatch.setattr(
        parsing,
        "census_names_one_row",
        lambda counts, totals, floor=1: original(counts, totals),
    )
    folder = tmp_path / "line"
    folder.mkdir(parents=True)
    table = folder / "table.csv"
    table.write_text(_judged_pair(2), encoding="utf-8", newline="")
    assert (
        _quiet(
            [
                "profile", f"{table}", "--out-dir", f"{folder}", "--replace",
                "--smallest-group", _FLOOR_ELEVEN, "--measurement", "reading",
            ]
        )
        == 0
    )
    block = _reading_block(_document(folder / "table-profile.json"))
    assert block["missing_by_source"]["-999"] == 20
    assert block["sentinel_verdicts"][0]["n_occurrences"] == 22
    monkeypatch.undo()
    assert (
        _quiet(
            [
                "generate", f"{folder / 'table-profile.json'}",
                "--out-dir", f"{folder}", "--seed", "4", "--replace",
            ]
        )
        == 1
    )
